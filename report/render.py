# Author: Shubhodeep Chatterjee
"""Render one report description into IEEE-style LaTeX (IEEEtran) and DOCX.

A report is a list of blocks:
  ("h1", text) ("h2", text) ("p", text)
  ("eq", latex, plain)                      displayed equation
  ("list", [text, ...])                     bulleted list
  ("table", label, caption, header, rows, colspec, wide)
  ("figure", label, path, caption, wide)
Inline text: plain Unicode. [@key] or [@a,b] cites, {ref:label} cross-references,
**bold**, _italic_. Unicode maths symbols are converted for LaTeX.
"""
import re
from pathlib import Path

UNI = {"θ": r"$\theta$", "λ": r"$\lambda$", "ε": r"$\varepsilon$", "≥": r"$\geq$", "≤": r"$\leq$",
       "×": r"$\times$", "→": r"$\rightarrow$", "≈": r"$\approx$", "±": r"$\pm$", "−": "-", "–": "--",
       "—": "---", "≠": r"$\neq$", "α": r"$\alpha$", "Σ": r"$\Sigma$", "∏": r"$\prod$", "·": r"$\cdot$", "ⁿ": r"$^n$", "Θ": r"$\Theta$", "∈": r"$\in$",
       "′": "'", "’": "'", "“": "``", "”": "''", "%": r"\%", "&": r"\&", "#": r"\#", "_": r"\_"}
_CITE = re.compile(r"\[@([^\]]+)\]")
_REF = re.compile(r"\{ref:([^}]+)\}")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_SUB = re.compile(r"([θΘλε])_(\w+)")
_GREEK = {"θ": r"\theta", "Θ": r"\Theta", "λ": r"\lambda", "ε": r"\varepsilon"}
_ITAL = re.compile(r"(?<![\w\\])_(.+?)_(?!\w)")


def tex_escape(s):
    out = []
    for ch in s:
        if ch in UNI:
            out.append(UNI[ch])
        elif ch in "{}":
            out.append("\\" + ch)
        elif ch == "~":
            out.append(r"\textasciitilde{}")
        elif ch == "^":
            out.append(r"\textasciicircum{}")
        elif ch == "<":
            out.append(r"$<$")
        elif ch == ">":
            out.append(r"$>$")
        elif ch == "\\":
            out.append(r"\textbackslash{}")
        else:
            out.append(ch)
    return "".join(out)


def tex_inline(s):
    # protect markup, escape the rest, restore
    tokens = []

    def keep(x):
        tokens.append(x)
        return f"\x00{len(tokens) - 1}\x00"
    s = _SUB.sub(lambda m: keep("$" + _GREEK[m.group(1)] + "_{" + m.group(2) + "}$"), s)
    s = _CITE.sub(lambda m: keep(r"\cite{" + ",".join(k.strip() for k in m.group(1).replace("@", "").split(",")) + "}"), s)
    s = _REF.sub(lambda m: keep(("Table~" if m.group(1).startswith("tab:") else "Fig.~") + r"\ref{" + m.group(1) + "}"), s)
    s = _BOLD.sub(lambda m: keep(r"\textbf{" + tex_escape(m.group(1)) + "}"), s)
    s = _ITAL.sub(lambda m: keep(r"\emph{" + tex_escape(m.group(1)) + "}"), s)
    s = tex_escape(s)
    return re.sub("\x00(\\d+)\x00", lambda m: tokens[int(m.group(1))], s)


def cite_order(meta, blocks):
    """Reference keys in order of first citation (IEEE numbering)."""
    order = []
    texts = [meta["abstract"]]
    for b in blocks:
        if b[0] == "p":
            texts.append(b[1])
        elif b[0] == "list":
            texts += b[1]
        elif b[0] == "table":
            texts.append(b[2])
        elif b[0] == "figure":
            texts.append(b[3])
    for s in texts:
        for m in _CITE.finditer(s):
            for k in m.group(1).replace("@", "").split(","):
                if k.strip() not in order:
                    order.append(k.strip())
    return order


def to_latex(meta, blocks, refs, out_path):
    L = [r"\documentclass[conference]{IEEEtran}",
         r"\IEEEoverridecommandlockouts",
         r"\usepackage{cite}", r"\usepackage{amsmath,amssymb}", r"\usepackage{graphicx}",
         r"\usepackage{booktabs}", r"\usepackage{array}", r"\usepackage[hidelinks]{hyperref}",
         r"\usepackage{xcolor}",
         r"\hypersetup{pdfauthor={" + meta["author"] + "}, pdftitle={" + meta["title"] + "}}",
         r"\graphicspath{{figures/}}",
         r"\begin{document}",
         r"\title{" + tex_escape(meta["title"]) + "}",
         r"\author{\IEEEauthorblockN{" + tex_escape(meta["author"]) + "}"
         + r"\IEEEauthorblockA{" + r" \\ ".join(tex_escape(x) for x in meta["affiliation"]) + "}}",
         r"\maketitle",
         r"\begin{abstract}", tex_inline(meta["abstract"]), r"\end{abstract}",
         r"\begin{IEEEkeywords}", tex_escape(meta["keywords"]), r"\end{IEEEkeywords}", ""]
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            L.append(r"\section{" + tex_escape(b[1]) + "}")
        elif kind == "h2":
            L.append(r"\subsection{" + tex_escape(b[1]) + "}")
        elif kind == "p":
            L.append(tex_inline(b[1]) + "\n")
        elif kind == "eq":
            L.append(r"\begin{equation}" + b[1] + r"\end{equation}")
        elif kind == "list":
            L.append(r"\begin{itemize}")
            L += [r"\item " + tex_inline(x) for x in b[1]]
            L.append(r"\end{itemize}")
        elif kind == "table":
            _, label, caption, header, rows, colspec, wide = b
            env = "table*" if wide else "table"
            L += [r"\begin{" + env + "}[!t]", r"\centering", r"\caption{" + tex_inline(caption) + "}",
                  r"\label{" + label + "}", r"\footnotesize", r"\setlength{\tabcolsep}{3pt}",
                  r"\begin{tabular}{" + colspec + "}", r"\toprule",
                  " & ".join(r"\textbf{" + tex_inline(h) + "}" for h in header) + r" \\", r"\midrule"]
            L += [" & ".join(tex_inline(str(c)) for c in r) + r" \\" for r in rows]
            L += [r"\bottomrule", r"\end{tabular}", r"\end{" + env + "}"]
        elif kind == "figure":
            _, label, path, caption, wide = b
            env = "figure*" if wide else "figure"
            width = r"\textwidth" if wide else r"\columnwidth"
            L += [r"\begin{" + env + "}[!t]", r"\centering",
                  r"\includegraphics[width=" + width + "]{" + Path(path).stem + "}",
                  r"\caption{" + tex_inline(caption) + "}", r"\label{" + label + "}", r"\end{" + env + "}"]
    L += [r"\begin{thebibliography}{99}"]
    refmap = dict(refs)
    order = cite_order(meta, blocks)
    missing = [k for k in order if k not in refmap]
    if missing:
        raise KeyError(f"cited but not in references: {missing}")
    for key in order:
        L.append(r"\bibitem{" + key + "} " + tex_escape(refmap[key]))
    L += [r"\end{thebibliography}", r"\end{document}"]
    Path(out_path).write_text("\n".join(L) + "\n")


# ---------------------------------------------------------------- DOCX
def to_docx(meta, blocks, refs, out_path):
    from docx import Document
    from docx.shared import Pt, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_SECTION
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    order = cite_order(meta, blocks)   # IEEE numbering by first citation
    labels = {}
    nt = nf = 0
    for b in blocks:
        if b[0] == "table":
            nt += 1; labels[b[1]] = _roman(nt)
        elif b[0] == "figure":
            nf += 1; labels[b[1]] = str(nf)
    num = {k: i + 1 for i, k in enumerate(order)}
    refmap = dict(refs)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = st.paragraph_format
    pf.space_after = Pt(0); pf.space_before = Pt(0); pf.first_line_indent = Cm(0.35)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.625)
    sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(1)

    def runs(par, s, size=None, base_bold=False, base_italic=False):
        s = _CITE.sub(lambda m: "[" + ", ".join(str(num[k.strip()]) for k in m.group(1).replace("@", "").split(",")) + "]", s)
        s = _REF.sub(lambda m: ("Table " if m.group(1).startswith("tab:") else "Fig. ") + labels.get(m.group(1), "?"), s)
        parts = re.split(r"(\*\*.+?\*\*|(?<![\w\\])_.+?_(?!\w))", s)
        for p in parts:
            if not p:
                continue
            bold = p.startswith("**") and p.endswith("**")
            ital = (not bold) and p.startswith("_") and p.endswith("_") and len(p) > 2
            txt = p[2:-2] if bold else (p[1:-1] if ital else p)
            for k, piece in enumerate(_SUB.split(txt)):
                # split yields text, greek, subscript, text, ...
                if not piece:
                    continue
                r = par.add_run(piece)
                r.bold = bold or base_bold
                r.italic = ital or base_italic or (k % 3 == 1)
                r.font.subscript = (k % 3 == 2)
                if size:
                    r.font.size = Pt(size)
        return par

    def para(s="", align=None, size=None, bold=False, italic=False, indent=True, space_after=0):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if not indent:
            p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(space_after)
        runs(p, s, size, bold, italic)
        return p

    para(meta["title"], WD_ALIGN_PARAGRAPH.CENTER, 22, indent=False, space_after=8)
    para(meta["author"], WD_ALIGN_PARAGRAPH.CENTER, 11, indent=False)
    for a in meta["affiliation"]:
        para(a, WD_ALIGN_PARAGRAPH.CENTER, 9, italic=True, indent=False)
    para("", indent=False, space_after=6)

    def set_columns(n):
        """Start a continuous section with n text columns."""
        sct = doc.add_section(WD_SECTION.CONTINUOUS)
        cols = sct._sectPr.xpath("./w:cols")
        c = cols[0] if cols else OxmlElement("w:cols")
        c.set(qn("w:num"), str(n)); c.set(qn("w:space"), str(int(0.25 * 1440)))
        if not cols:
            sct._sectPr.append(c)

    set_columns(2)                     # two-column body

    p = para("", indent=False, space_after=4)
    r = p.add_run("Abstract—"); r.bold = True; r.italic = True; r.font.size = Pt(9)
    runs(p, meta["abstract"], 9, base_bold=True)
    p = para("", indent=False, space_after=8)
    r = p.add_run("Index Terms—"); r.bold = True; r.italic = True; r.font.size = Pt(9)
    runs(p, meta["keywords"], 9, base_bold=True, base_italic=True)

    h1n = 0; h2n = 0
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            h1n += 1; h2n = 0
            para(f"{_roman(h1n)}. {b[1].upper()}", WD_ALIGN_PARAGRAPH.CENTER, 10, indent=False, space_after=3).paragraph_format.space_before = Pt(8)
        elif kind == "h2":
            h2n += 1
            q = para(f"{chr(64 + h2n)}. {b[1]}", WD_ALIGN_PARAGRAPH.LEFT, 10, italic=True, indent=False, space_after=2)
            q.paragraph_format.space_before = Pt(4)
        elif kind == "p":
            para(b[1])
        elif kind == "eq":
            para(b[2], WD_ALIGN_PARAGRAPH.CENTER, italic=True, indent=False, space_after=3)
        elif kind == "list":
            for x in b[1]:
                q = para("• " + x, indent=False)
                q.paragraph_format.left_indent = Cm(0.4)
        elif kind == "table":
            _, label, caption, header, rows, colspec, wide = b
            if wide:
                set_columns(1)
            para(f"TABLE {labels[label]}", WD_ALIGN_PARAGRAPH.CENTER, 8, indent=False).paragraph_format.space_before = Pt(6)
            para(caption.upper() if len(caption) < 60 else caption, WD_ALIGN_PARAGRAPH.CENTER, 8, indent=False, space_after=2)
            t = doc.add_table(rows=1 + len(rows), cols=len(header))
            t.style = "Table Grid"
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]; cell.text = ""
                runs(cell.paragraphs[0], h, 7, base_bold=True)
            for i, row in enumerate(rows):
                for j, v in enumerate(row):
                    cell = t.rows[i + 1].cells[j]; cell.text = ""
                    runs(cell.paragraphs[0], str(v), 7)
            # column widths proportional to content, filling the text width
            total = 7.25 if wide else 3.4
            lens = [max([len(str(h))] + [len(str(r[j])) for r in rows]) for j, h in enumerate(header)]
            lens = [min(max(l, 4), 60) + 3 for l in lens]       # +3: cell padding
            t.autofit = False
            widths = [Inches(total * l / sum(lens)) for l in lens]
            for j, w in enumerate(widths):
                for row in t.rows:
                    row.cells[j].width = w
            grid = t._tbl.tblGrid                      # LibreOffice reads the grid, Word the cells
            for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
                gc.set(qn("w:w"), str(int(w.inches * 1440)))
            tblPr = t._tbl.tblPr
            tw = tblPr.find(qn("w:tblW"))
            if tw is None:
                tw = OxmlElement("w:tblW"); tblPr.append(tw)
            tw.set(qn("w:w"), str(int(total * 1440))); tw.set(qn("w:type"), "dxa")
            lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
            for row in t.rows:
                for cell in row.cells:
                    for q in cell.paragraphs:
                        q.paragraph_format.first_line_indent = Cm(0)
                        q.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para("", indent=False, space_after=4)
            if wide:
                set_columns(2)
        elif kind == "figure":
            _, label, path, caption, wide = b
            if wide:
                set_columns(1)
            q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.first_line_indent = Cm(0)
            q.add_run().add_picture(str(path), width=Inches(6.5 if wide else 3.4))
            para(f"Fig. {labels[label]}. " + caption, WD_ALIGN_PARAGRAPH.JUSTIFY, 8, indent=False, space_after=6)
            if wide:
                set_columns(2)
    para("REFERENCES", WD_ALIGN_PARAGRAPH.CENTER, 10, indent=False, space_after=3).paragraph_format.space_before = Pt(8)
    for k in order:
        q = para(f"[{num[k]}] " + refmap[k], size=8, indent=False)
        q.paragraph_format.left_indent = Cm(0.5); q.paragraph_format.first_line_indent = Cm(-0.5)
    cp = doc.core_properties
    cp.author = meta["author"]; cp.last_modified_by = meta["author"]; cp.title = meta["title"]
    cp.subject = "Cytogenetic progression models from the Mitelman Database"; cp.keywords = meta["keywords"]
    cp.comments = ""
    import datetime as _dt
    cp.created = cp.modified = _dt.datetime.now().replace(microsecond=0)
    doc.save(out_path)


def _roman(n):
    vals = [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s; n -= v
    return out
